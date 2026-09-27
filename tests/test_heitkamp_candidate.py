"""Industrial sampling is not independent biological replication or methane data."""

import importlib.util
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location(
    "heitkamp", ROOT / "scripts/build_heitkamp_2021_candidate.py"
)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def workbook():
    w = Workbook()
    s = w.active
    s.title = "pH"
    s.cell(1, 2, "pH")
    for n in range(7):
        row = 3 + n * 4
        s.cell(row, 2, f"BGP{n + 1}")
        s.cell(row + 1, 2, "Day of sampling [d]")
        s.cell(row + 1, 3, 0)
        s.cell(row + 2, 2, "pH")
        s.cell(row + 2, 3, 7)
    return w


def test_sparse_cells_and_text_sentinels_are_not_zero():
    w = workbook()
    s = w.active
    s["C5"] = None
    s["C9"] = "n.a"
    s["C13"] = 0
    rows, qc = module.extract(w, "fixture")
    assert rows[0]["value_source"] is None
    assert rows[1]["value_source"] == "n.a"
    assert rows[2]["value_source"] == 0
    assert qc["missing"] == qc["nonnumeric"] == qc["zeros"] == 1


@pytest.mark.parametrize("fault", ["duplicate", "orphan_value"])
def test_ambiguous_times_fail_closed(fault):
    w = workbook()
    w.active["D5"] = 8
    if fault == "duplicate":
        w.active["D4"] = 0
    with pytest.raises(ValueError):
        module.extract(w, "fixture")


def test_modified_source_is_rejected(tmp_path):
    path = tmp_path / "source.xlsx"
    path.write_bytes(b"not the audited workbook")
    with pytest.raises(ValueError, match="SHA-256"):
        module.build(path)


def test_public_candidate_preserves_scope_and_unit_conflicts():
    path = ROOT / "results/intake"
    rows = pd.read_csv(path / "heitkamp_2021_chemistry.csv", keep_default_na=False)
    qc = json.loads((path / "heitkamp_2021_qc.json").read_text())
    assert len(rows) == qc["measurement_rows"] == 1365
    assert rows.plant_id.nunique() == 7
    assert rows[["plant_id", "sampling_day"]].drop_duplicates().shape[0] == 91
    assert rows.analyte.nunique() == 15
    assert rows.source_value_cell.nunique() == 1365
    assert (rows.value_source == "n.a").sum() == 10
    assert (rows.value_source == "0").sum() == 402
    assert not rows.model_admission.any()
    assert rows.loc[rows.analyte == "VS", "qc_flags"].str.contains("basis_conflict").all()
    assert "methane_productivity_validation" in qc["forbidden_uses"]
