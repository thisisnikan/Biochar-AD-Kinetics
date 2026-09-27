"""Build a quarantined, source-traceable Sanglier methane candidate table.

This is an intake/QC artifact, not a training or external-validation dataset.
The source workbook is public but is downloaded separately and verified by MD5.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SOURCE_DOI = "10.57745/BUJORT"
SOURCE_MD5 = "282f5076857dd4f8ca38089a6d499909"
SOURCE_FILE_ID = 152903
KEY = ("Lab", "Batch", "Bottle")
OUTPUT_COLUMNS = (
    "study_id",
    "source_doi",
    "source_md5",
    "lab",
    "experiment_id",
    "batch_id",
    "bottle_id",
    "replicate",
    "source_condition_ampts",
    "source_condition_inoculation",
    "biochar_label_ampts",
    "inoculation_biochar_g",
    "te_supplementation_ampts",
    "te_solution_ampts",
    "inoculation_te1_ml",
    "inoculation_te2_ml",
    "inoculation_ammonium_bicarbonate_g",
    "inoculation_substrate_vs_g",
    "inoculation_date",
    "absolute_days_lab",
    "batch_start_days_lab",
    "within_batch_days",
    "volume_raw_nml",
    "volume_corrected_nml",
    "flow_nml_h",
    "methane_yield_nml_gvs",
    "source_amptshort_row",
    "source_inoculation_row",
    "qc_flags",
)


def sheet_records(workbook, name: str) -> list[dict]:
    sheet = workbook[name]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if len(header) != len(set(header)) or any(v is None for v in header):
        raise ValueError(f"Invalid or duplicate header in {name}")
    return [
        dict(zip(header, row, strict=True), source_row=n)
        for n, row in enumerate(rows, 2)
        if any(v is not None for v in row)
    ]


def source_hash(path: Path, expected_md5: str = SOURCE_MD5) -> str:
    digest = hashlib.md5()  # Checks the deposit's published checksum, not a security boundary.
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    if digest.hexdigest() != expected_md5:
        raise ValueError("Process.xlsx MD5 differs from the published deposit checksum")
    return digest.hexdigest()


def _value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def build(source: Path, *, expected_md5: str = SOURCE_MD5):
    md5 = source_hash(source, expected_md5)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        inventory = {
            s.title: {
                "rows_including_header": s.max_row,
                "columns": s.max_column,
                "state": s.sheet_state,
            }
            for s in workbook
        }
        inoculations = sheet_records(workbook, "Inoculation")
        methane = sheet_records(workbook, "AMPTShort")
    finally:
        workbook.close()

    inoculation_by_key = {}
    for record in inoculations:
        if any(record.get(k) is None for k in KEY):
            continue  # ancillary formula/calculation row, not a bottle record
        key = tuple(record[k] for k in KEY)
        if key in inoculation_by_key:
            raise ValueError(f"Duplicate Inoculation bottle/batch key: {key}")
        inoculation_by_key[key] = record

    starts = {}
    seen = set()
    for record in methane:
        key = tuple(record.get(k) for k in (*KEY, "Days"))
        if any(v is None for v in key):
            raise ValueError(
                f"Incomplete methane identity/time at AMPTShort!{record['source_row']}"
            )
        if key in seen:
            raise ValueError(f"Duplicate methane bottle/batch/time: {key}")
        seen.add(key)
        if key[:3] not in inoculation_by_key:
            raise ValueError(f"No Inoculation record for AMPTShort!{record['source_row']}")
        cycle = key[:2]
        starts[cycle] = min(float(record["Days"]), starts.get(cycle, float("inf")))

    output = []
    flagged_cycles = {}
    for record in methane:
        key = tuple(record[k] for k in KEY)
        inoc = inoculation_by_key[key]
        flags = []
        biochar_label = record["Biochar"]  # source labels 0/1/2; no g/L conversion
        biochar_g = inoc["Biochar"]
        if biochar_label == 0 and biochar_g is not None and biochar_g > 0:
            flags.append("control_label_positive_inoculation_biochar")
        if str(record["Cond"]).replace(" ", "") != str(inoc["Cond"]).replace(" ", ""):
            flags.append("condition_label_disagreement")
        if record["Replicate"] != inoc["Replicate"]:
            flags.append("replicate_disagreement")
        if record["Experiment_ID"] != inoc["Experiment_ID"]:
            flags.append("experiment_id_disagreement")
        if flags:
            flagged_cycles[key] = flags
        start = starts[key[:2]]
        output.append(
            dict(
                zip(
                    OUTPUT_COLUMNS,
                    map(
                        _value,
                        (
                            "sanglier_2022",
                            SOURCE_DOI,
                            md5,
                            record["Lab"],
                            record["Experiment_ID"],
                            record["Batch"],
                            record["Bottle"],
                            record["Replicate"],
                            record["Cond"],
                            inoc["Cond"],
                            biochar_label,
                            biochar_g,
                            record["TE_supplementation"],
                            record["TE_solution"],
                            inoc["TE 1"],
                            inoc["TE 2"],
                            inoc["Ammonium bicarbonate"],
                            inoc["VS Substrate"],
                            inoc["Date"],
                            record["Days"],
                            start,
                            round(float(record["Days"]) - start, 10),
                            record["Volume_raw"],
                            record["Volume"],
                            record["Flow"],
                            record["Methane yield"],
                            f"AMPTShort!{record['source_row']}",
                            f"Inoculation!{inoc['source_row']}",
                            ";".join(flags),
                        ),
                    ),
                    strict=True,
                )
            )
        )

    counts = Counter(flag for row in output for flag in row["qc_flags"].split(";") if flag)
    report = {
        "status": "CANDIDATE_QC_ONLY_NOT_ADMITTED_FOR_FITTING_OR_VALIDATION",
        "source_doi": SOURCE_DOI,
        "source_file_id": SOURCE_FILE_ID,
        "source_md5": md5,
        "sheet_inventory": inventory,
        "amptshort_rows": len(methane),
        "inoculation_bottle_batch_rows": len(inoculation_by_key),
        "joined_rows": len(output),
        "duplicate_methane_keys": 0,
        "qc_flag_row_counts": dict(sorted(counts.items())),
        "flagged_bottle_batches": [
            {"lab": k[0], "batch": k[1], "bottle": k[2], "flags": v}
            for k, v in sorted(flagged_cycles.items())
        ],
        "missing_counts": {
            c: sum(row[c] is None for row in output)
            for c in (
                "volume_raw_nml",
                "volume_corrected_nml",
                "flow_nml_h",
                "methane_yield_nml_gvs",
            )
        },
        "zero_counts": {
            c: sum(row[c] == 0 for row in output)
            for c in (
                "volume_raw_nml",
                "volume_corrected_nml",
                "flow_nml_h",
                "methane_yield_nml_gvs",
            )
        },
        "notes": [
            "AMPTShort is the source's 12-hour derivative of AMPTS, not independent observations.",
            "Days is absolute within each lab; within_batch_days subtracts that lab/batch's minimum.",
            "Numeric zero and missing cells are preserved distinctly; no blank correction is applied.",
            "Analyses and Events sheets are inventoried but not joined; VFA/TAN validation is pending.",
            "Biochar labels 0/1/2 and inoculation biochar grams are separate source fields; no g/L dose is inferred.",
            "Trace-element and other co-interventions preclude biochar-only causal attribution.",
        ],
    }
    return output, report


def write(rows, report, output: Path, qc_report: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    qc_report.parent.mkdir(parents=True, exist_ok=True)
    # Gzip has no varying filename or timestamp, making rebuilds byte-stable.
    with (
        output.open("wb") as raw,
        gzip.GzipFile(fileobj=raw, mode="wb", mtime=0, filename="") as zipped,
        io.TextIOWrapper(zipped, encoding="utf-8", newline="") as text,
    ):
        writer = csv.DictWriter(text, fieldnames=OUTPUT_COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    qc_report.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_xlsx", type=Path)
    parser.add_argument(
        "--output", type=Path, default=Path("results/intake/sanglier_2022_candidate.csv.gz")
    )
    parser.add_argument(
        "--qc-report", type=Path, default=Path("results/intake/sanglier_2022_qc.json")
    )
    args = parser.parse_args()
    write(*build(args.source_xlsx), args.output, args.qc_report)


if __name__ == "__main__":
    main()
