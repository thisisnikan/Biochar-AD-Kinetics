"""Extract industrial digester chemistry without pooling plants or inferring gas yield."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SOURCE_DOI = "10.1186/s13068-021-02034-5"
SOURCE_URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1186%2Fs13068-021-02034-5/MediaObjects/"
    "13068_2021_2034_MOESM1_ESM.xlsx"
)
SOURCE_SHA256 = "c60a45ab498d3fbbb25ece27d6cd338b348742f8f03ca8bd126d4a5ffd0ded00"
STATUS = "SOURCE_VERIFIED_CHEMISTRY_CANDIDATE_NOT_ADMITTED_FOR_VALIDATION"


def extract(workbook, sha256):
    output, seen = [], set()
    for sheet in workbook:
        rows = list(sheet.iter_rows())
        blocks = [
            i
            for i, r in enumerate(rows)
            if len(r) > 1 and isinstance(r[1].value, str) and r[1].value.startswith("BGP")
        ]
        if len(blocks) != 7:
            raise ValueError(f"Expected seven plant blocks in {sheet.title}")
        for i in blocks:
            plant = rows[i][1].value
            days, values = rows[i + 1], rows[i + 2]
            if days[1].value != "Day of sampling [d]":
                raise ValueError("Unexpected source time label")
            for time, value in zip(days[2:], values[2:], strict=True):
                if time.value is None and value.value is None:
                    continue
                if not isinstance(time.value, (float, int)) or time.value < 0:
                    raise ValueError("Value without valid sampling day")
                key = (plant, sheet.title, time.value)
                if key in seen:
                    raise ValueError(f"Duplicate plant/analyte/day {key}")
                seen.add(key)
                flags = []
                if value.value is None:
                    flags.append("missing_value")
                elif not isinstance(value.value, (float, int)):
                    flags.append("nonnumeric_value_requires_review")
                if sheet.title == "VS":
                    flags.append("workbook_fresh_biomass_vs_article_TS_basis_conflict")
                if sheet.title == "TVFA":
                    flags.append("source_total_basis_not_recomputed")
                output.append(
                    {
                        "study_id": "heitkamp_2021",
                        "source_doi": SOURCE_DOI,
                        "source_sha256": sha256,
                        "plant_id": plant,
                        "sampling_day": time.value,
                        "analyte": sheet.title,
                        "value_source": value.value,
                        "source_header": rows[0][1].value,
                        "source_value_label": values[1].value,
                        "source_time_cell": f"{sheet.title}!{time.coordinate}",
                        "source_value_cell": f"{sheet.title}!{value.coordinate}",
                        "qc_flags": ";".join(flags),
                        "model_admission": False,
                    }
                )
    counts = Counter(r["analyte"] for r in output)
    report = {
        "status": STATUS,
        "source_doi": SOURCE_DOI,
        "source_url": SOURCE_URL,
        "source_sha256": sha256,
        "source_license": "CC-BY-4.0",
        "license_url": "https://creativecommons.org/licenses/by/4.0/",
        "measurement_rows": len(output),
        "plants": sorted({r["plant_id"] for r in output}),
        "plant_sampling_days": len({(r["plant_id"], r["sampling_day"]) for r in output}),
        "counts_by_analyte": dict(sorted(counts.items())),
        "zeros": sum(r["value_source"] == 0 for r in output),
        "missing": sum(r["value_source"] is None for r in output),
        "nonnumeric": sum(
            r["value_source"] is not None and not isinstance(r["value_source"], (int, float))
            for r in output
        ),
        "min_sampling_day": min(r["sampling_day"] for r in output),
        "max_sampling_day": max(r["sampling_day"] for r in output),
        "forbidden_uses": [
            "methane_productivity_validation",
            "causal_biochar_effect",
            "treat_timepoints_as_independent_reactor_replicates",
        ],
        "notes": [
            "Wide worksheet blocks transposed without resampling, imputation or normalization.",
            "Workbook values may already aggregate measurements; no replicate identity invented.",
            "Source unit labels retained; NH4-N kg/t is not converted to g/L.",
            "VS source header says fresh biomass; article figure caption says percent TS.",
            "TVFA source totals retained; analytical-equivalent basis not independently verified.",
            "Day zero is the workbook origin, not an independently verified dosing timestamp.",
        ],
    }
    return output, report


def build(source):
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("Heitkamp source SHA-256 differs from the audited publisher workbook")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return extract(workbook, digest)
    finally:
        workbook.close()


def write(rows, report, destination):
    destination.mkdir(parents=True, exist_ok=True)
    with (destination / "heitkamp_2021_chemistry.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    (destination / "heitkamp_2021_qc.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_xlsx", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("results/intake"))
    args = parser.parse_args()
    write(*build(args.source_xlsx), args.output_dir)


if __name__ == "__main__":
    main()
