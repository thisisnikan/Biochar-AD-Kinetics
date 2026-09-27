"""Ingest Marta's private Frontiers workbook as summary data, never raw replicates.

Outputs default to ignored results/private/. The second feeding is kept separate;
its recorded day-23 zero conflicts with the article's day-22 feeding origin.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SOURCE_SHA256 = "c731306efea9f4fc8cfa77c5b4c65e68450212e3364ccc134c6dee6d5d3f9355"
SOURCE_DOI = "10.3389/fceng.2024.1384495"
LABELS = [
    "Cellulose",
    "Control",
    "B1-1",
    "B1-5",
    "B1-10",
    "B2-1",
    "B2-5",
    "B2-10",
    "B3-1",
    "B3-5",
    "B3-10",
]
DAYS = [0, 2, 5, 7, 12, 15, 22, 23, 26, 29, 33, 40, 44]


def numeric_or_missing(value):
    if value is None or value == "-":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TypeError(f"Unexpected numeric value: {value!r}")
    if not math.isfinite(value) or value < 0:
        raise ValueError("Negative/nonfinite methane mean or standard deviation")
    return value


def transform(sheet, design):
    if list(sheet[9][2:15]) != DAYS or list(sheet[24][2:15]) != DAYS:
        raise ValueError("Mean/SD time grids differ from the audited workbook")
    output = []
    for offset, label in enumerate(LABELS):
        mean_row, sd_row = 10 + offset, 25 + offset
        if sheet[mean_row][1] != label or sheet[sd_row][1] != label:
            raise ValueError("Mean/SD treatment alignment differs")
        condition = (
            label.lower()
            if label in ("Control", "Cellulose")
            else ("bc" + label[1:].replace("-", "_") + "pct")
        )
        if condition not in design:
            raise ValueError(f"Unknown condition {condition}")
        metadata = design[condition]
        for index, day in enumerate(DAYS, 2):
            raw_mean, raw_sd = sheet[mean_row][index], sheet[sd_row][index]
            mean, sd = numeric_or_missing(raw_mean), numeric_or_missing(raw_sd)
            if (mean is None) != (sd is None):
                raise ValueError("Mean and SD missingness differs")
            feeding = 1 if day <= 22 else 2
            unavailable = feeding == 2 and label in ("Cellulose", "B1-1", "B1-5", "B1-10")
            if unavailable != (mean is None):
                raise ValueError("Missingness differs from the reported feeding design")
            flags = ["summary_only_no_reactor_trajectories"]
            if feeding == 2:
                flags.append("day23_zero_vs_article_day22_feeding_origin_unresolved")
            if unavailable:
                flags.append("not_selected_for_second_feeding")
            if mean == 0 and sd == 0:
                flags.append("reported_zero_sd_not_infinite_precision")
            col = get_column_letter(index + 1)
            output.append(
                {
                    "study_id": "garcia_prats_2024",
                    "source_doi": SOURCE_DOI,
                    "source_sha256": SOURCE_SHA256,
                    "source_label": label,
                    "condition_id": condition,
                    "biochar_id": metadata["biochar_id"],
                    "condition_type": metadata["condition_type"],
                    "dose_pct_ts": metadata["dose_pct_ts"],
                    "dose_basis": "substrate_plus_inoculum_TS",
                    "nominal_dose_g_l": metadata["nominal_dose_g_l"],
                    "dose_conversion_basis": "published_mass_mg_divided_by_working_volume_ml",
                    "temperature_c": metadata["temperature_c"],
                    "feeding_id": feeding,
                    "substrate_batch_id": f"OFMSW_S{feeding}"
                    if label != "Cellulose"
                    else "cellulose",
                    "source_day": day,
                    "days_since_recorded_phase_zero": day if feeding == 1 else day - 23,
                    "phase_time_is_confirmed": feeding == 1,
                    "cumulative_methane_mean_ml_gvs": mean,
                    "reported_sd_ml_gvs": sd,
                    "n_replicates_reported": 3,
                    "replicate_level_available": False,
                    "measurement_available": mean is not None,
                    "source_mean_token": raw_mean,
                    "source_sd_token": raw_sd,
                    "source_mean_cell": f"Hoja1!{col}{mean_row + 1}",
                    "source_sd_cell": f"Hoja1!{col}{sd_row + 1}",
                    "source_time_cell": f"Hoja1!{col}10",
                    "source_sd_time_cell": f"Hoja1!{col}25",
                    "model_admission": False,
                    "qc_flags": ";".join(flags),
                }
            )
    report = {
        "status": "PRIVATE_SUMMARY_INTAKE_NOT_REACTOR_LEVEL_VALIDATION",
        "source_sha256": SOURCE_SHA256,
        "source_doi": SOURCE_DOI,
        "source_filename": "Data Frontiers.xlsx",
        "source_received_date": "2026-09-18",
        "redistribution_permission": "not_established",
        "rows": len(output),
        "observed_rows": sum(r["measurement_available"] for r in output),
        "missing_design_slots": sum(not r["measurement_available"] for r in output),
        "conditions": len(LABELS),
        "source_is_summary": True,
        "individual_replicates_available": False,
        "observations_by_feeding": {
            str(f): sum(r["feeding_id"] == f and r["measurement_available"] for r in output)
            for f in (1, 2)
        },
        "blockers": [
            "second_feeding_time_origin_conflict",
            "individual_reactor_data_unavailable",
            "normalization_and_blank_correction_details_need_confirmation",
        ],
        "restrictions": [
            "no_synthetic_replicates",
            "no_cross_feeding_causal_comparison",
            "no_shared_control_duplication_as_independent_evidence",
            "no_inverse_variance_weights_from_zero_SD",
            "no_raw_workbook_or_derived_rows_in_public_git",
        ],
    }
    return output, report


def build(source, design_path):
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("Source SHA-256 differs from the author-shared workbook")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Hoja1"]:
            raise ValueError("Unexpected workbook sheets")
        sheet = list(workbook["Hoja1"].values)
    finally:
        workbook.close()
    with design_path.open(newline="") as f:
        design = {r["condition_id"]: r for r in csv.DictReader(f)}
    rows, report = transform(sheet, design)
    report["design_sha256"] = hashlib.sha256(design_path.read_bytes()).hexdigest()
    return rows, report


def write(rows, report, destination):
    destination.mkdir(parents=True, exist_ok=True)
    with (destination / "methane_summary.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    (destination / "intake_qc.json").write_text(json.dumps(report, indent=2, sort_keys=True) + "\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument(
        "--design",
        type=Path,
        default=Path("data/experimental/garcia_prats_2024_treatment_design.csv"),
    )
    parser.add_argument(
        "--output-dir", type=Path, default=Path("results/private/garcia_prats_2024")
    )
    args = parser.parse_args()
    write(*build(args.source, args.design), args.output_dir)


if __name__ == "__main__":
    main()
